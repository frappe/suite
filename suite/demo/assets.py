"""Generate small, self-contained Northwind spreadsheets, decks, and uploaded files."""

import json
from html import escape
from io import BytesIO
from pathlib import Path

import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from suite import drive

# PNGs are 256 px renders of the SVG sources beside them; re-render after editing one.
BRAND = Path(__file__).parent / "brand"
BUDGET = [
    ["Q4 2026 operating budget", None, None, None, None],
    ["Category", "October", "November", "December", "Quarter"],
    ["People", 42000, 44000, 44000, "=SUM(B3:D3)"],
    ["Infrastructure", 3200, 3600, 4000, "=SUM(B4:D4)"],
    ["Research incentives", 900, 1200, 900, "=SUM(B5:D5)"],
    ["Launch marketing", 800, 2500, 1200, "=SUM(B6:D6)"],
    ["Travel and equipment", 1800, 2400, 1600, "=SUM(B7:D7)"],
    ["Total", "=SUM(B3:B7)", "=SUM(C3:C7)", "=SUM(D3:D7)", "=SUM(E3:E7)"],
    ["Monthly average", "=AVERAGE(B8:D8)"],
]
HIRING = [
    ["Q4 hiring tracker", None, None, None, None],
    ["Role", "Team", "Stage", "Annual salary", "Interview score"],
    ["Backend engineer", "Engineering", "Technical interview", 96000, 4.2],
    ["Product designer", "Design", "Portfolio review", 88000, 4.5],
    ["Customer researcher", "Product", "Screening", 72000, 3.8],
    ["Customer success lead", "Marketing", "Offer", 78000, 4.6],
    ["Planned salary total", None, None, "=SUM(D3:D6)"],
    ["Average interview score", None, None, None, "=AVERAGE(E3:E6)"],
]
PIPELINE = [
    ["November sales pipeline", None, None, None, None],
    ["Account", "Stage", "Annual value", "Probability", "Weighted value"],
    ["Cedar Health", "Pilot", 18000, 0.6, "=C3*D3"],
    ["Atlas Robotics", "Proposal", 24000, 0.4, "=C4*D4"],
    ["Harbor Studio", "Negotiation", 9600, 0.8, "=C5*D5"],
    ["Fieldwork Co", "Discovery", 14400, 0.2, "=C6*D6"],
    ["Total", None, "=SUM(C3:C6)", None, "=SUM(E3:E6)"],
    ["Average deal", None, "=AVERAGE(C3:C6)"],
]
LAUNCH = [
    ("Beacon collections", "Private beta | 2 November 2026", "Evidence beside every decision"),
    (
        "The customer problem",
        "Research is scattered across projects.\nRoadmap reviews start with a search for evidence.",
        "20 minutes lost per review",
    ),
    (
        "A shared collection",
        "Collect notes without making copies.\nInvite readers, commenters, and editors.",
        "One place to discuss the evidence",
    ),
    (
        "Our beta plan",
        "12 customer teams, two weeks.\nWeekly office hours and a recorded walkthrough.",
        "Maya: product | Theo: invitations",
    ),
    (
        "How we measure success",
        "40% create a second collection in 14 days.\nTrack time to first share and seven-day return rate.",
        "Review on 16 November",
    ),
]
REVIEW = [
    ("Q3 review / Q4 priorities", "Northwind Labs | October 2026", "A smaller scope, a stronger product"),
    (
        "Customer momentum",
        "32 paying teams, up from 24 in June.\nWeekly active teams grew 28%.",
        "91% quarterly retention",
    ),
    (
        "Quality and reliability",
        "Import success improved from 94% to 99.2%.\nTwo incidents; both recovered in under 30 minutes.",
        "Next: mixed-date CSV coverage",
    ),
    (
        "Where we invested",
        "Research: 18 customer interviews.\nDelivery: faster search and accessible sharing.",
        "$156k quarterly operating spend",
    ),
    (
        "Q4 commitments",
        "Ship collections beta.\nHire a backend engineer and a product designer.",
        "Protect focus: no additional product lines",
    ),
]


def spreadsheet(parent: str, imports: str, title: str, rows: list[list], kind: str) -> str:
    """Import a styled XLSX and retain the original as an ordinary uploaded file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(row)
    sheet.merge_cells("A1:E1")
    sheet.freeze_panes = "B3"
    sheet.auto_filter.ref = f"A2:E{len(rows)}"
    for row in sheet:
        for cell in row:
            if cell.row <= 2:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="234D47")
            elif cell.column > 1 and (isinstance(cell.value, int | float) or cell.data_type == "f"):
                cell.number_format = '"$"#,##0.00'
                if kind == "pipeline" and cell.column == 4:
                    cell.number_format = "0%"
                if kind == "hiring" and cell.column == 5:
                    cell.number_format = "0.0"
    for column in range(1, 6):
        sheet.column_dimensions[get_column_letter(column)].width = 30 if column == 1 else 24
    notes = workbook.create_sheet("Read me")
    notes.append(["Northwind Labs - internal planning"])
    notes.append(["All amounts are USD. Figures and candidates are fictional."])
    notes.append(["Owner: Sofia Martinez" if kind != "pipeline" else "Owner: Theo Williams"])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    source = drive.store_file(imports, f"{title}.xlsx", stream)
    return drive.import_document(parent, title, content_doctype="Sheet", from_node=source)


def presentation(parent: str, title: str, story: list[tuple[str, str, str]]) -> str:
    node = drive.create_document(parent, title, content_doctype="Presentation")
    deck = frappe.get_doc("Presentation", frappe.db.get_value("Presentation", {"node": node}))
    deck.slides = []
    for number, (heading, body, callout) in enumerate(story, 1):
        elements = [
            _shape(f"accent-{number}", 56, 50, 66, 6, "#5FE0B5FF"),
            _text(f"heading-{number}", heading, 56, 100, 800, 44),
            _text(f"body-{number}", body, 56, 214, 800, 26),
            _shape(f"card-{number}", 56, 380, 848, 76, "#234D47FF"),
            _text(f"callout-{number}", callout, 78, 398, 800, 22),
            _text(f"footer-{number}", f"NORTHWIND LABS     /     {number:02}", 56, 492, 800, 13),
        ]
        deck.append(
            "slides",
            {
                "background": "#102D29FF",
                "elements": json.dumps(elements),
                "transition": "None",
                "transition_duration": "0",
            },
        )
    deck.save(ignore_permissions=True)
    return node


def uploaded_files() -> list[tuple[str, bytes]]:
    """Answer real files without writing temporary assets into the working tree."""
    image = Image.new("RGB", (720, 420))
    image.putdata([(15 + x // 12, 45 + y // 7, 65 + (x + y) // 14) for y in range(420) for x in range(720)])
    drawing = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=34)
    drawing.ellipse((490, 40, 670, 220), fill="#5fe0b5")
    drawing.text((40, 285), "NORTHWIND LABS", fill="white", font=font)
    drawing.text(
        (42, 340), "Evidence beside every decision", fill="#b5e5d9", font=ImageFont.load_default(size=22)
    )
    cover = _image_bytes(image, "PNG")
    chart = Image.new("RGB", (720, 420), "#f5f8f7")
    drawing = ImageDraw.Draw(chart)
    drawing.text((35, 25), "Weekly active teams", fill="#234d47", font=font)
    for index, (label, value) in enumerate((("Jul", 18), ("Aug", 23), ("Sep", 29), ("Oct", 32))):
        left = 80 + index * 155
        drawing.rectangle((left, 340 - value * 7, left + 85, 340), fill="#248c73")
        drawing.text((left, 355), f"{label}: {value}", fill="#234d47", font=ImageFont.load_default(size=20))
    page = Image.new("RGB", (850, 1100), "white")
    drawing = ImageDraw.Draw(page)
    lines = [
        "Northwind Labs",
        "Equipment and remote work policy",
        "Effective: 1 October 2026",
        "",
        "Equipment allowance: USD 1,500 on joining.",
        "Annual home office allowance: USD 400.",
        "Submit itemized receipts within 30 days.",
        "Manager approval is required above USD 250.",
        "",
        "Security: use device encryption and a screen lock.",
        "Report a lost device to Arjun immediately.",
        "",
        "Questions? Contact Sofia in Finance.",
        "Fictional company policy for product testing.",
    ]
    for index, line in enumerate(lines):
        drawing.text((60, 70 + index * 58), line, fill="#234d47", font=ImageFont.load_default(size=24))
    return [
        ("Northwind brand cover.png", cover),
        ("Weekly active teams.png", _image_bytes(chart, "PNG")),
        ("Equipment policy.pdf", _image_bytes(page, "PDF")),
        (
            "Beta customer feedback.csv",
            b"account,theme,priority\nCedar Health,Sharing clarity,High\nAtlas Robotics,CSV dates,High\nHarbor Studio,Keyboard navigation,Medium\n",
        ),
        (
            "Launch checklist.md",
            b"# Beta launch\n\n- [x] Approve copy\n- [ ] Verify imports\n- [ ] Invite twelve customer teams\n\nOwner: Theo. Review on 29 October.\n",
        ),
        (
            "Read me.txt",
            b"Northwind Labs demo workspace\nAll people, customers, and financial figures are fictional.\nStart with Company / Handbook / Welcome to Northwind.\n",
        ),
    ]


def brand_file(source: str, file_name: str, doctype: str, name: str, field: str) -> str:
    """Attach a stored brand image as a public file and answer its URL."""
    file = frappe.get_doc(
        {
            "doctype": "File",
            "file_name": file_name,
            "content": (BRAND / source).read_bytes(),
            "is_private": 0,
            "attached_to_doctype": doctype,
            "attached_to_name": name,
            "attached_to_field": field,
        }
    ).insert(ignore_permissions=True)
    return file.file_url


def _image_bytes(image: Image.Image, kind: str) -> bytes:
    stream = BytesIO()
    image.save(stream, format=kind)
    return stream.getvalue()


def _text(identifier: str, text: str, left: int, top: int, width: int, size: int) -> dict:
    paragraphs = "".join(
        f'<p><span style="font-family: Inter; font-size: {size}px; color: #FFFFFFFF;">{escape(line)}</span></p>'
        for line in text.splitlines()
    )
    return {
        "id": identifier,
        "type": "text",
        "left": left,
        "top": top,
        "width": width,
        "content": paragraphs,
        "opacity": 100,
        "rotation": 0,
        "editorMetadata": {"lineHeight": 1.4},
    }


def _shape(identifier: str, left: int, top: int, width: int, height: int, color: str) -> dict:
    return {
        "id": identifier,
        "type": "shape",
        "shapeType": "rectangle",
        "left": left,
        "top": top,
        "width": width,
        "height": height,
        "fillColor": color,
        "strokeWidth": 0,
        "strokeColor": color,
        "strokeStyle": "solid",
        "borderRadius": 8,
        "opacity": 100,
        "rotation": 0,
        "shadowBlur": 0,
        "shadowOffset": 0,
    }
