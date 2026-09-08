# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and Contributors
# See license.txt
"""Sheets' adoption of the Drive content contract (ticket 19, §6.7, §10.7, §14.6).

Adoption was an expand phase. Sheets declared its `ContentTypeSpec` at ticket
19 and `Sheet` gained the `node` Link; ticket 28 linked every row, and ticket
29 made the registry entry and the hook changes together. So `suite/hooks.py`
now names `suite.sheets.drive.SPEC` and points `Sheet`, `Sheet Op Log`, and
`Sheet Collab State` at `suite.drive.framework`. `Sheet Snapshot` keeps its own
guard: §14.6 migrates its rows into `Drive Node Version`, so it stays a Build
source until Cleanup.

The classes here:

`TestSheetsDeclaration`   the declaration itself, and the hook entries the app
                          ships. No rows and no database.
`TestSheetsWorkbook`      the xlsx importer, the version envelope, and the media
                          scan, exercised as the pure functions they are.
`TestSheetsAfterActivation`
                          what activation settled: a sheet needs its node, and
                          every legacy path refuses a linked one.
`TestSheetsInDrive`       the Drive-native lifecycle, versions, purge, media,
                          and satellites.

**A sheet with no node cannot exist any more.** `require_node` holds §5.13 for
a registered doctype (`content.py:757-770`), so the legacy-sheet cases this
module used to carry are gone with the state they described. `activated()`
stays as a name so every call site reads the same, and it is now only the
registry cache drop.

The first two classes are plain `unittest.TestCase`: they mock `frappe.db` where
they need it and run without a site. The last two need real rows.
"""

from __future__ import annotations

import dataclasses
import datetime
import inspect
import io
import json
import pathlib
import unittest
import zipfile
from contextlib import contextmanager
from unittest import mock

import frappe
from frappe.tests import IntegrationTestCase

from suite import drive
from suite import hooks as suite_hooks
from suite.drive._core.access import grant
from suite.drive._core.content import clear_registry_cache, governs, spec_for
from suite.drive._core.errors import DriveConflict, DriveForbidden, DriveNotFound
from suite.drive._core.nodes import create_file, create_folder, purge, update
from suite.drive._core.principals import Principals
from suite.drive._core.roots import create_root, purge_root, update_root
from suite.drive._core.versions import restore_version
from suite.drive.framework import satellite_has_permission, satellite_query_conditions
from suite.sheets import drive as sheets
from suite.sheets.doctype.sheet.storage import decode_sheets_data, encode_sheets_data
from suite.tests.utils import ensure_user

USER = "sheets-adoption-user@example.com"
OTHER = "sheets-adoption-other@example.com"

DOCTYPE = "Sheet"
OP_LOG = "Sheet Op Log"
COLLAB_STATE = "Sheet Collab State"


@contextmanager
def activated():
    """Read the registry `suite/hooks.py` ships, and leave nothing behind.

    Ticket 29 registered `suite.sheets.drive.SPEC` and moved the permission
    hooks with it, so there is nothing to inject. The registry is built from
    `drive_content_types` and cached per request, and the cache is dropped on
    the way in and on the way out.
    """
    clear_registry_cache()
    try:
        yield
    finally:
        clear_registry_cache()


# ── xlsx fixtures ────────────────────────────────────────────────────────────


def workbook_bytes(build) -> bytes:
    """Render one openpyxl workbook `build` filled in, as xlsx bytes."""
    from openpyxl import Workbook

    book = Workbook()
    build(book)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def sheet_doctype_json() -> dict:
    """The shipped `Sheet` DocType JSON, read from disk rather than a site."""
    path = pathlib.Path(sheets.__file__).parent / "doctype" / "sheet" / "sheet.json"
    return json.loads(path.read_text())


def _doc_perm(role: str) -> dict:
    for row in sheet_doctype_json()["permissions"]:
        if row.get("role") == role:
            return row
    raise AssertionError(f"no DocPerm row for role {role}")


def simple_workbook() -> bytes:
    def build(book):
        sheet = book.active
        sheet.title = "Data"
        sheet["A1"] = "Name"
        sheet["B1"] = "Amount"
        sheet["A2"] = "Widget"
        sheet["B2"] = 42.0
        sheet["B3"] = "=SUM(B2:B2)"

    return workbook_bytes(build)


def respliced_workbook(**parts: bytes) -> bytes:
    """A real xlsx package with named members replaced verbatim.

    openpyxl writes a valid package; this swaps one part for a hostile one, so
    every test below starts from a file Excel would open.
    """
    book = workbook_bytes(lambda book: book.active.__setitem__("A1", "x"))
    source = zipfile.ZipFile(io.BytesIO(book))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for item in source.namelist():
            package.writestr(item, parts.get(item, source.read(item)))
    return output.getvalue()


def _reordered_tabs(raw: bytes, order: list[str]) -> bytes:
    """Rewrite `xl/workbook.xml` so the tab order differs from the part order.

    Excel does this whenever a tab is dragged: the `<sheet>` elements move and
    `sheet1.xml` stays where it is.
    """
    from xml.etree import ElementTree

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ElementTree.register_namespace("", namespace)
    source = zipfile.ZipFile(io.BytesIO(raw))
    tree = ElementTree.fromstring(source.read("xl/workbook.xml"))
    sheets_element = tree.find(f"{{{namespace}}}sheets")
    by_name = {element.get("name"): element for element in list(sheets_element)}
    for element in list(sheets_element):
        sheets_element.remove(element)
    for name in order:
        sheets_element.append(by_name[name])
    rewritten = ElementTree.tostring(tree, encoding="UTF-8", xml_declaration=True)

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for item in source.namelist():
            package.writestr(item, rewritten if item == "xl/workbook.xml" else source.read(item))
    return output.getvalue()


def worksheet_xml(body: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"{body}</worksheet>"
    ).encode()


# ── the declaration ──────────────────────────────────────────────────────────


class TestSheetsDeclaration(unittest.TestCase):
    """What §10.7 fixes about Sheets, and the hook entries ticket 29 installed."""

    def test_it_declares_the_identity_section_ten_seven_fixes(self):
        self.assertEqual(sheets.SPEC.doctype, DOCTYPE)
        self.assertEqual(sheets.SPEC.mime, "frappe/sheet")
        self.assertEqual(sheets.SPEC.node_field, "node")

    def test_a_sheet_has_no_default_export_and_no_export_at_all(self):
        """§10.7: Sheets stays hidden over WebDAV for this release."""
        self.assertIsNone(sheets.SPEC.default_export)
        self.assertIsNone(sheets.SPEC.export)
        self.assertEqual(sheets.SPEC.export_formats, ())

    def test_drive_never_renders_a_sheet_and_sheets_never_pushes_one(self):
        self.assertFalse(sheets.SPEC.pushes_preview)

    def test_the_three_required_callbacks_are_declared(self):
        self.assertIs(sheets.SPEC.create_empty, sheets.create_empty)
        self.assertIs(sheets.SPEC.duplicate, sheets.duplicate)
        self.assertIs(sheets.SPEC.on_purge, sheets.on_purge)

    def test_an_xlsx_can_become_a_sheet(self):
        self.assertIs(sheets.SPEC.import_from_file, sheets.import_from_file)

    def test_a_sheet_has_versions_it_can_be_restored_from(self):
        self.assertIs(sheets.SPEC.version_bytes, sheets.version_bytes)
        self.assertIs(sheets.SPEC.restore_version, sheets.restore_version)

    def test_the_two_satellites_are_the_op_log_and_the_collaborative_document(self):
        self.assertEqual(
            [(one.doctype, one.link_field) for one in sheets.SPEC.satellites],
            [(OP_LOG, "sheet"), (COLLAB_STATE, "sheet")],
        )

    def test_sheet_snapshot_is_not_a_satellite_because_build_still_reads_it(self):
        """§14.6 migrates its rows into `Drive Node Version` (§14.10 drops it)."""
        self.assertNotIn("Sheet Snapshot", [one.doctype for one in sheets.SPEC.satellites])

    def test_the_media_sweep_can_ask_what_a_body_names(self):
        self.assertIs(sheets.SPEC.used_nodes, sheets.used_nodes)

    def test_no_remap_media_because_a_body_names_no_media_yet(self):
        """Declaring a no-op would let a copy carry media nothing names (§8.9)."""
        self.assertIsNone(sheets.SPEC.remap_media)

    def test_the_four_cleanup_pending_columns_are_declared(self):
        self.assertEqual(sheets.SPEC.legacy_fields, ("title", "trashed", "trashed_on", "trashed_by"))

    def test_head_snapshot_needs_no_exemption_because_ten_two_allows_it(self):
        from suite.drive._core.content import FORBIDDEN_FIELD_NAMES

        self.assertNotIn("head_snapshot", FORBIDDEN_FIELD_NAMES)

    # the `All` DocPerm, which the guards deny against

    def test_the_open_baseline_row_still_carries_every_legacy_right(self):
        """§10.4 needs the row open; it does not need it narrower than before.

        A Frappe permission hook can only deny, so the owner rule comes from
        the MANAGE grant Build wrote on the owner's own node, and
        `doc_has_permission` reads it. A right this row drops is a right no
        hook can return: `share_sheet` asks `ptype="share"` and
        `frappe.share.check_share_permission` asks it again
        (`frappe/share.py:239`), so dropping `share` here refuses the owner.
        """
        baseline = _doc_perm("All")
        for right in ("read", "write", "create", "delete", "share", "export", "print", "email", "report"):
            self.assertEqual(baseline.get(right), 1, f"the `All` row must keep `{right}`")

    def test_the_open_baseline_row_no_longer_restricts_itself_to_the_owner(self):
        self.assertNotIn("if_owner", _doc_perm("All"))

    def test_guest_reads_only_through_a_link_grant(self):
        guest = _doc_perm("Guest")
        self.assertEqual(guest.get("read"), 1)
        for right in ("write", "create", "delete", "share"):
            self.assertNotIn(right, guest)

    # the hooks ticket 29 moved

    # Read from `suite.hooks` rather than `frappe.get_hooks`: the module is what
    # the site ships, and reading it directly means these assertions hold with
    # no site bound.

    def test_the_declaration_is_registered(self):
        """README execution rules: stage the registry after the node links
        exist. Ticket 28 wrote them and `refuse_unlinked_documents` is what
        holds the ordering on a real migration."""
        self.assertIn("suite.sheets.drive.SPEC", suite_hooks.drive_content_types)

    def test_both_sheet_permission_hooks_are_now_the_framework_s(self):
        self.assertEqual(suite_hooks.has_permission[DOCTYPE], "suite.drive.framework.doc_has_permission")
        self.assertEqual(
            suite_hooks.permission_query_conditions[DOCTYPE],
            "suite.drive.framework.doc_query_conditions",
        )

    def test_both_satellites_take_the_framework_s_satellite_hooks(self):
        """`Sheet Collab State` had no entry at all before ticket 29, so it
        gains both rather than moving them."""
        for doctype in (OP_LOG, COLLAB_STATE):
            with self.subTest(doctype=doctype):
                self.assertEqual(
                    suite_hooks.has_permission[doctype],
                    "suite.drive.framework.satellite_has_permission",
                )
                self.assertEqual(
                    suite_hooks.permission_query_conditions[doctype],
                    "suite.drive.framework.satellite_query_conditions",
                )

    def test_sheet_snapshot_keeps_its_own_guard(self):
        """No spec declares it and its role rows are open, so the guard is the
        only thing holding the migrated history back. §14.6 reads the rows into
        `Drive Node Version` and §14.10 drops the doctype with the guard."""
        self.assertEqual(
            suite_hooks.has_permission["Sheet Snapshot"],
            "suite.sheets.permissions.sheet_snapshot_has_permission",
        )
        self.assertEqual(
            suite_hooks.permission_query_conditions["Sheet Snapshot"],
            "suite.sheets.permissions.sheet_snapshot_query",
        )

    def test_the_registry_answers_with_exactly_this_declaration(self):
        with activated():
            self.assertTrue(governs(DOCTYPE))
            self.assertIs(spec_for(DOCTYPE), sheets.SPEC)

    def test_the_declaration_is_frozen_so_nothing_can_edit_it_at_runtime(self):
        with self.assertRaises(dataclasses.FrozenInstanceError):
            sheets.SPEC.default_export = "xlsx"

    def test_the_drive_entry_points_keep_the_argument_order_this_module_calls(self):
        """Pin the leading positionals, because a swap here is silent.

        `nodes` and `versions` take `Principals` first; `grant` takes it
        fourth; the two framework hooks take `user` first because Frappe calls
        them that way (`frappe/model/db_query.py:1343`). Passing a node where
        `Principals` belongs reaches the database as a filter and fails deep
        inside the query builder, far from the call that made the mistake.
        """
        expected = {
            create_file: ["principals", "parent", "title"],
            create_folder: ["principals", "parent", "title"],
            update: ["principals", "node"],
            purge: ["principals", "node"],
            restore_version: ["principals", "node", "seq"],
            grant: ["node_id", "principal", "role", "principals"],
            satellite_query_conditions: ["user", "doctype"],
            satellite_has_permission: ["doc", "ptype", "user", "debug"],
        }
        for function, leading in expected.items():
            with self.subTest(function=function.__name__):
                parameters = list(inspect.signature(function).parameters)
                self.assertEqual(parameters[: len(leading)], leading)


# ── the body, as pure functions ──────────────────────────────────────────────


class TestSheetsWorkbook(unittest.TestCase):
    """The importer, the version envelope, and the media scan. No database."""

    def _frappe(self, **db):
        """Replace `suite.sheets.drive.frappe` — nothing here reads a row.

        `frappe.db` is an unbound proxy outside a request, so the module's own
        reference is what gets replaced, exactly as the other Sheets unit tests
        do it.
        """
        patcher = mock.patch("suite.sheets.drive.frappe")
        patched = patcher.start()
        self.addCleanup(patcher.stop)
        patched.ValidationError = ValueError
        patched.DoesNotExistError = LookupError
        patched.throw.side_effect = lambda msg, exc=Exception: (_ for _ in ()).throw(exc(msg))
        for name, value in db.items():
            getattr(patched.db, name).return_value = value
        return patched

    # the version envelope

    def test_version_bytes_carries_the_workbook_and_the_head_seq(self):
        self._frappe(get_value=frappe._dict(sheets_data=None, head_seq=7))
        stream, mime = sheets.version_bytes("SH-1")
        self.assertEqual(mime, "application/json")
        payload = json.loads(stream.read())
        self.assertEqual(payload["schema"], "sheet/1")
        self.assertEqual(payload["head_seq"], 7)
        self.assertEqual(payload["sheets_data"], "{}")

    def test_version_bytes_decodes_the_stored_envelope(self):
        stored = encode_sheets_data(json.dumps({"sheet": {"v": 2}}))
        self._frappe(get_value=frappe._dict(sheets_data=stored, head_seq=1))
        stream, _mime = sheets.version_bytes("SH-1")
        payload = json.loads(stream.read())
        self.assertEqual(json.loads(payload["sheets_data"]), {"sheet": {"v": 2}})

    def test_a_payload_that_is_not_a_sheet_envelope_is_refused_whole(self):
        """A truncated workbook would read as an empty one, so it is refused."""
        self._frappe()
        for raw in (
            b"",
            b"not json",
            json.dumps({"sheets_data": "{}"}).encode(),
            json.dumps({"schema": "deck/1", "sheets_data": "{}"}).encode(),
        ):
            with self.subTest(raw=raw), self.assertRaises(Exception):
                sheets._version_payload(raw)

    def test_a_payload_whose_workbook_is_not_json_is_refused(self):
        self._frappe()
        raw = json.dumps({"schema": "sheet/1", "sheets_data": "{oops"}).encode()
        with self.assertRaises(Exception):
            sheets._version_payload(raw)

    def test_a_valid_payload_survives_the_round_trip(self):
        self._frappe()
        raw = json.dumps({"schema": "sheet/1", "sheets_data": '{"a":1}', "head_seq": 3}).encode()
        self.assertEqual(sheets._version_payload(raw), {"sheets_data": '{"a":1}', "head_seq": 3})

    def test_a_missing_head_seq_reads_as_zero_rather_than_failing(self):
        self._frappe()
        raw = json.dumps({"schema": "sheet/1", "sheets_data": "{}"}).encode()
        self.assertEqual(sheets._version_payload(raw)["head_seq"], 0)

    # the media scan

    def _used(self, body: dict | str) -> set[str]:
        stored = encode_sheets_data(body if isinstance(body, str) else json.dumps(body))
        self._frappe(get_value=stored)
        return sheets.used_nodes("SH-1")

    def test_it_reports_every_id_shaped_token_at_every_depth(self):
        found = self._used({"sheet": {"sheets": {"Data": {"rows": {"0": ["nodeone", {"src": "nodetwo"}]}}}}})
        self.assertIn("nodeone", found)
        self.assertIn("nodetwo", found)

    def test_it_over_reports_rather_than_letting_the_sweep_trash_live_media(self):
        """§10.6 trashes what the app does not report, so a plain word costs nothing."""
        self.assertIn("total", self._used({"sheet": {"sheets": {"Data": {"rows": {"0": ["total"]}}}}}))

    def test_a_token_no_node_id_can_be_is_not_reported(self):
        found = self._used({"cells": ["=SUM(A1:A2)", "https://example.com/x", "a b", ""]})
        self.assertNotIn("=SUM(A1:A2)", found)
        self.assertNotIn("https://example.com/x", found)
        self.assertNotIn("a b", found)

    def test_an_unreadable_body_reports_what_it_can_read_instead_of_raising(self):
        """One bad row must not stop the daily pass."""
        self._frappe(get_value="{oops nodeone")
        self.assertIn("nodeone", sheets.used_nodes("SH-1"))

    def test_a_sheet_that_is_gone_names_nothing(self):
        self._frappe(get_value=None, exists=False)
        self.assertEqual(sheets.used_nodes("SH-1"), set())

    # the xlsx importer

    def _import(self, raw: bytes) -> dict:
        self._frappe()
        return json.loads(sheets._workbook_from_xlsx(raw))

    def test_it_writes_the_version_two_pack_the_client_reads(self):
        packed = self._import(simple_workbook())["sheet"]
        self.assertEqual(packed["v"], 2)
        self.assertEqual(packed["current"], "Data")
        self.assertEqual(list(packed["sheets"]), ["Data"])

    def test_cells_arrive_row_major_at_their_own_columns(self):
        rows = self._import(simple_workbook())["sheet"]["sheets"]["Data"]["rows"]
        self.assertEqual(rows["0"], ["Name", "Amount"])
        self.assertEqual(rows["1"], ["Widget", "42"])

    def test_a_formula_keeps_its_leading_equals(self):
        rows = self._import(simple_workbook())["sheet"]["sheets"]["Data"]["rows"]
        self.assertEqual(rows["2"], [None, "=SUM(B2:B2)"])

    def test_every_worksheet_becomes_a_sub_sheet(self):
        def build(book):
            book.active.title = "First"
            book.active["A1"] = "a"
            second = book.create_sheet("Second")
            second["A1"] = "b"

        packed = self._import(workbook_bytes(build))["sheet"]
        self.assertEqual(list(packed["sheets"]), ["First", "Second"])
        self.assertEqual(packed["current"], "First")

    def test_a_whole_float_arrives_as_an_integer_string_like_the_browser(self):
        def build(book):
            book.active["A1"] = 42.0
            book.active["A2"] = 42.5

        rows = self._import(workbook_bytes(build))["sheet"]["sheets"]["Sheet"]["rows"]
        self.assertEqual(rows["0"], ["42"])
        self.assertEqual(rows["1"], ["42.5"])

    def test_a_boolean_arrives_as_the_engine_spells_it(self):
        def build(book):
            book.active["A1"] = True
            book.active["A2"] = False

        rows = self._import(workbook_bytes(build))["sheet"]["sheets"]["Sheet"]["rows"]
        self.assertEqual([rows["0"][0], rows["1"][0]], ["TRUE", "FALSE"])

    def test_a_date_arrives_as_the_iso_text_the_engine_reads_back(self):
        def build(book):
            book.active["A1"] = datetime.datetime(2026, 9, 6)
            book.active["A2"] = datetime.datetime(2026, 9, 6, 13, 30, 5)

        rows = self._import(workbook_bytes(build))["sheet"]["sheets"]["Sheet"]["rows"]
        self.assertEqual(rows["0"], ["2026-09-06"])
        self.assertEqual(rows["1"], ["2026-09-06 13:30:05"])

    def test_merges_come_across_as_the_master_and_slave_maps(self):
        def build(book):
            sheet = book.active
            sheet.title = "Data"
            sheet["A1"] = "banner"
            sheet.merge_cells("A1:C2")

        merged = self._import(workbook_bytes(build))["merge"]["Data"]
        self.assertEqual(merged["masterMap"], {"A1": {"rowSpan": 2, "colSpan": 3, "r": 0, "c": 0}})
        self.assertEqual(sorted(merged["slaveMap"]), ["A2", "B1", "B2", "C1", "C2"])
        self.assertEqual(set(merged["slaveMap"].values()), {"A1"})

    def test_a_workbook_with_no_merges_writes_null_like_the_client_does(self):
        self.assertIsNone(self._import(simple_workbook())["merge"])

    def test_merges_stay_with_their_own_worksheet(self):
        def build(book):
            first = book.active
            first.title = "First"
            first["A1"] = "x"
            second = book.create_sheet("Second")
            second["A1"] = "y"
            second.merge_cells("A1:B1")

        merged = self._import(workbook_bytes(build))["merge"]
        self.assertEqual(list(merged), ["Second"])

    def test_a_number_format_travels_as_the_engine_s_format_string(self):
        def build(book):
            sheet = book.active
            sheet.title = "Data"
            sheet["A1"] = 0.125
            sheet["A1"].number_format = "0.00%"
            sheet["A2"] = datetime.datetime(2026, 9, 6)
            sheet["A2"].number_format = "yyyy-mm-dd"
            sheet["A3"] = "text"
            sheet["A3"].number_format = "@"

        cells = self._import(workbook_bytes(build))["formats"]["Data"]["cells"]
        self.assertEqual(cells["A1"], {"numberFormat": "percentage"})
        self.assertEqual(cells["A2"], {"numberFormat": "date:ymd"})
        self.assertEqual(cells["A3"], {"numberFormat": "text"})

    def test_a_general_format_is_not_recorded_at_all(self):
        self.assertNotIn("Data", self._import(simple_workbook()).get("formats", {}))

    def test_an_unknown_format_code_is_preserved_verbatim(self):
        def build(book):
            book.active.title = "Data"
            book.active["A1"] = 1
            book.active["A1"].number_format = '"€"#,##0.00'

        cells = self._import(workbook_bytes(build))["formats"]["Data"]["cells"]
        self.assertEqual(cells["A1"], {"numberFormat": 'custom:"€"#,##0.00'})

    def test_a_file_openpyxl_cannot_read_is_refused_with_a_plain_message(self):
        self._frappe()
        with self.assertRaises(sheets.UnreadableWorkbook):
            sheets._workbook_from_xlsx(b"this is not a spreadsheet")

    def test_a_zip_that_is_not_a_workbook_is_refused_too(self):
        self._frappe()
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            package.writestr("hello.txt", "hi")
        with self.assertRaises(sheets.UnreadableWorkbook):
            sheets._workbook_from_xlsx(buffer.getvalue())

    # the format table, in isolation

    def test_the_number_format_table_maps_both_ways_the_client_does(self):
        cases = {
            "": "",
            "General": "",
            "@": "text",
            "dd/mm/yyyy": "date:dmy",
            "hh:mm:ss": "time:hms",
            "yyyy-mm-dd hh:mm": "datetime:ymd_hm",
            "0.00%": "percentage",
            "0.0%": "percentage:1",
            "#,##0": "number",
            "#,##0.00": "number:2",
        }
        for code, expected in cases.items():
            with self.subTest(code=code):
                self.assertEqual(sheets._number_format(code), expected)

    def test_a_quoted_or_escaped_percent_does_not_scale_the_value(self):
        self.assertTrue(sheets._number_format('#,##0"%"').startswith("custom:"))
        self.assertTrue(sheets._number_format("#,##0\\%").startswith("custom:"))

    # what a hostile workbook costs

    def test_one_merge_range_cannot_expand_into_the_whole_grid(self):
        """`A1:XFD1048576` is legal, is 4 KB on the wire, and covers 17e9 cells.

        Before the bound, `_merge_slice` wrote one `slaveMap` entry for each of
        them: measured, that took a worker past 2 GB in 20 seconds.
        """
        evil = respliced_workbook(
            **{
                "xl/worksheets/sheet1.xml": worksheet_xml(
                    '<sheetData><row r="1"><c r="A1" t="inlineStr">'
                    "<is><t>x</t></is></c></row></sheetData>"
                    '<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>'
                )
            }
        )
        self._frappe()
        with self.assertRaises(ValueError) as refusal:
            sheets._workbook_from_xlsx(evil)
        self.assertIn("merges more than", str(refusal.exception))

    def test_an_ordinary_merge_still_comes_across(self):
        """The bound must refuse the bomb and nothing a person would send."""

        def build(book):
            book.active.title = "Data"
            book.active["A1"] = "x"
            book.active.merge_cells("A1:C3")

        merge = self._import(workbook_bytes(build))["merge"]["Data"]
        self.assertEqual(merge["masterMap"]["A1"]["rowSpan"], 3)
        self.assertEqual(merge["slaveMap"]["C3"], "A1")

    def test_a_sparse_row_cannot_allocate_the_columns_it_skips(self):
        """A row is a dense list, so one cell at XFD costs 16384 slots.

        `MAX_IMPORT_CELLS` counts values, not slots, so 20000 cells in a 105 KB
        file allocated 327 million of them: measured, past 2 GB in 8 seconds.
        """
        rows = "".join(
            f'<row r="{r}"><c r="XFD{r}" t="inlineStr"><is><t>v</t></is></c></row>' for r in range(1, 501)
        )
        evil = respliced_workbook(
            **{"xl/worksheets/sheet1.xml": worksheet_xml(f"<sheetData>{rows}</sheetData>")}
        )
        self._frappe()
        with self.assertRaises(ValueError) as refusal:
            sheets._workbook_from_xlsx(evil)
        self.assertIn("too sparse", str(refusal.exception))

    def test_a_wide_but_honest_row_is_not_refused(self):
        def build(book):
            for column in range(1, 51):
                book.active.cell(row=1, column=column, value=column)

        rows = self._import(workbook_bytes(build))["sheet"]["sheets"]["Sheet"]["rows"]
        self.assertEqual(len(rows["0"]), 50)

    def test_a_package_that_claims_to_expand_past_the_bound_is_refused(self):
        """A zip bomb is small on the wire and large once inflated."""
        evil = respliced_workbook(
            **{"xl/worksheets/sheet1.xml": worksheet_xml("<sheetData/>") + b" " * (2 * 1024 * 1024)}
        )
        self._frappe()
        with mock.patch.object(sheets, "MAX_IMPORT_UNZIPPED", 64 * 1024):
            with self.assertRaises(ValueError) as refusal:
                sheets._workbook_from_xlsx(evil)
        self.assertIn("expands to more than", str(refusal.exception))

    def test_a_part_carrying_a_document_type_declaration_is_refused(self):
        """`xml.etree.ElementTree` expands internal entities, so a DTD is a bomb.

        A spreadsheet never carries one, so its presence is the whole check.
        """
        bomb = (
            b'<?xml version="1.0"?>'
            b'<!DOCTYPE worksheet [<!ENTITY a "aaaaaaaaaa">'
            b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b"<sheetData/></worksheet>"
        )
        evil = respliced_workbook(**{"xl/worksheets/sheet1.xml": bomb})
        self._frappe()
        with self.assertRaises(sheets.UnreadableWorkbook):
            sheets._workbook_from_xlsx(evil)

    def test_a_merge_lands_on_its_own_worksheet_whatever_the_part_order(self):
        """Part numbering does not track tab order once a tab is moved.

        Pairing `sheetnames` with the parts sorted by suffix gave Beta's merge
        to Gamma. The workbook's own relationship is the authority.
        """

        def build(book):
            book.active.title = "Gamma"
            book.active["A1"] = "g"
            book.create_sheet("Alpha")["A1"] = "a"
            beta = book.create_sheet("Beta")
            beta["A1"] = "b"
            beta.merge_cells("A1:B1")

        raw = workbook_bytes(build)
        source = zipfile.ZipFile(io.BytesIO(raw))
        # Beta is the last tab and the last part, so the bug is only visible
        # once the two orders differ. Reorder the tabs, not the parts.
        merged = self._import(_reordered_tabs(raw, ["Beta", "Gamma", "Alpha"]))["merge"]
        self.assertIn("xl/worksheets/sheet3.xml", source.namelist())
        self.assertEqual(list(merged), ["Beta"])
        self.assertEqual(merged["Beta"]["slaveMap"]["B1"], "A1")

    def test_a_restore_drops_the_collaborative_document_it_replaces(self):
        """A live Y.Doc outranks `sheets_data`, so a restore must clear it."""
        patched = self._frappe(exists=True)
        version = json.dumps(
            {"schema": "sheet/1", "sheets_data": json.dumps({"sheet": {"v": 2}}), "head_seq": 3}
        ).encode()
        with mock.patch.object(sheets, "_append_op", return_value=9):
            sheets.restore_version("SH-1", io.BytesIO(version))
        patched.db.delete.assert_called_once_with("Sheet Collab State", {"sheet": "SH-1"})

    def test_a_workbook_that_declares_too_many_merge_ranges_is_refused(self):
        many = "".join(f'<mergeCell ref="A{r}:B{r}"/>' for r in range(1, 12))
        evil = respliced_workbook(
            **{
                "xl/worksheets/sheet1.xml": worksheet_xml(
                    '<sheetData><row r="1"><c r="A1" t="inlineStr">'
                    "<is><t>x</t></is></c></row></sheetData>"
                    f'<mergeCells count="11">{many}</mergeCells>'
                )
            }
        )
        self._frappe()
        with mock.patch.object(sheets, "MAX_IMPORT_MERGE_RANGES", 10):
            with self.assertRaises(ValueError) as refusal:
                sheets._workbook_from_xlsx(evil)
        self.assertIn("merged ranges", str(refusal.exception))

    def test_column_labels_match_the_client_and_the_cell_codec(self):
        from suite.sheets.doctype.sheet.cell_codec import _col_label

        for index in (0, 1, 25, 26, 27, 51, 52, 701, 702):
            with self.subTest(index=index):
                self.assertEqual(sheets._column_label(index), _col_label(index))
                self.assertEqual(sheets._column_index(sheets._column_label(index)), index)


# ── what activation settled ──────────────────────────────────────────────────


class TestSheetsAfterActivation(IntegrationTestCase):
    """A sheet needs its node, and every legacy path refuses a linked one.

    A sheet with no node cannot exist here: `require_node` holds §5.13 for a
    registered doctype (`content.py:757-770`), so the legacy-sheet cases this
    class used to carry describe a state Build and ticket 29 removed together.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        ensure_user(USER)
        _purge_fixture_roots()
        frappe.db.commit()
        cls.addClassCleanup(_drop_fixture_users)

    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        self.addCleanup(self._remove_fixture_rows)
        self.addCleanup(frappe.set_user, "Administrator")

    def _remove_fixture_rows(self):
        frappe.set_user("Administrator")
        _purge_fixture_roots()
        for name in frappe.get_all(DOCTYPE, filters={"title": ("like", "adoption-%")}, pluck="name"):
            frappe.db.delete(OP_LOG, {"sheet": name})
            frappe.db.delete("Sheet Seq", {"sheet": name})
            frappe.db.delete("DocShare", {"share_doctype": DOCTYPE, "share_name": name})
            frappe.delete_doc(DOCTYPE, name, force=1, ignore_permissions=True, ignore_missing=True)
            _drop_backing_file(name)
        frappe.db.commit()

    def _linked_sheet(self) -> tuple[str, str]:
        """One sheet linked the way Build links it: through Drive, then read back."""
        with activated():
            root = create_root(kind="Personal", title="Sheets Root", user=USER)
            node = drive.create_document(root.node, "adoption-linked", content_doctype=DOCTYPE)
        return node, frappe.db.get_value("Drive Node", node, "content_docname")

    # the row Build left behind cannot be written again

    def test_a_sheet_with_no_node_is_refused_on_insert(self):
        """The state every deleted case in this class described. Registration
        is what makes the node mandatory, and `DriveContent.before_insert` is
        where the refusal lands (`content.py:606-607`)."""
        with self.assertRaises(DriveConflict):
            frappe.get_doc({"doctype": DOCTYPE, "title": "adoption-orphan", "sheets_data": "{}"}).insert()

    # a linked row refuses every legacy path

    def test_a_linked_sheet_gets_no_backing_file(self):
        """A second backing row would be a second answer to "who may open this"."""
        from suite.drive.overrides.file import File as DriveFile

        _node, docname = self._linked_sheet()
        self.assertFalse(DriveFile.get_for_doc(DOCTYPE, docname))

    def test_every_legacy_endpoint_refuses_a_linked_sheet(self):
        from suite.sheets import api

        _node, docname = self._linked_sheet()
        calls = {
            "get_sheet_shares": lambda: api.get_sheet_shares(docname),
            "share_sheet": lambda: api.share_sheet(docname, OTHER, write=1),
            "unshare_sheet": lambda: api.unshare_sheet(docname, OTHER),
            "delete_sheet": lambda: api.delete_sheet(docname),
            "restore_sheet": lambda: api.restore_sheet(docname),
            "delete_sheet_permanent": lambda: api.delete_sheet_permanent(docname),
            "rename_sheet": lambda: api.rename_sheet(docname, "nope"),
            "duplicate_sheet": lambda: api.duplicate_sheet(docname),
        }
        for label, call in calls.items():
            with self.subTest(endpoint=label), self.assertRaises(frappe.ValidationError):
                call()

    def test_the_sheets_trash_cascade_refuses_a_linked_sheet(self):
        from suite.sheets.trash import hard_delete_sheet

        _node, docname = self._linked_sheet()
        with self.assertRaises(frappe.ValidationError):
            hard_delete_sheet(docname)

    def test_the_nightly_purge_never_sees_a_linked_sheet(self):
        from suite.sheets.trash import purge_trashed_sheets

        _node, docname = self._linked_sheet()
        # Even with the frozen column forced past the retention window.
        frappe.db.set_value(
            DOCTYPE,
            docname,
            {"trashed": 1, "trashed_on": "2020-01-01 00:00:00"},
            update_modified=False,
        )
        purge_trashed_sheets()
        self.assertTrue(frappe.db.exists(DOCTYPE, docname))


# ── the Drive-native lifecycle ───────────────────────────────────────────────


class TestSheetsInDrive(IntegrationTestCase):
    """Create, copy, import, version, restore, purge, media, and satellites.

    Every test runs under `activated()`, which is now the registry cache drop
    alone: ticket 29 registered the declaration these workflows read.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        ensure_user(USER)
        ensure_user(OTHER)
        _purge_fixture_roots()
        frappe.db.commit()
        cls.addClassCleanup(_drop_fixture_users)

    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        self._blobs_before = set(frappe.get_all("File Blob", pluck="name"))
        # Entered first, so its exit runs last: the fixture purge below is a
        # Drive workflow and reads the registry this drops the cache for.
        activation = activated()
        activation.__enter__()
        self.addCleanup(activation.__exit__, None, None, None)
        self.addCleanup(self._remove_fixture_rows)
        self.root = create_root(kind="Personal", title="Sheets Root", user=USER)
        self.other_root = create_root(kind="Personal", title="Sheets Other", user=OTHER)
        self.admin = Principals("Administrator", ("Administrator",), (), is_admin=True)
        self.person = Principals(USER, (USER, "$GENERAL"), ("$PUBLIC",))

    def _remove_fixture_rows(self):
        frappe.set_user("Administrator")
        _purge_fixture_roots()
        for blob in set(frappe.get_all("File Blob", pluck="name")) - self._blobs_before:
            frappe.delete_doc("File Blob", blob, force=1, ignore_permissions=True, ignore_missing=True)
        frappe.db.commit()

    # helpers

    def _sheet(self, title="Sheet", parent=None, **kwargs) -> str:
        return drive.create_document(parent or self.root.node, title, content_doctype=DOCTYPE, **kwargs)

    def _docname(self, node: str) -> str:
        return frappe.db.get_value("Drive Node", node, "content_docname")

    def _body(self, node: str) -> str:
        return decode_sheets_data(frappe.db.get_value(DOCTYPE, self._docname(node), "sheets_data"))

    def _write_body(self, node: str, body: dict) -> None:
        frappe.db.set_value(DOCTYPE, self._docname(node), "sheets_data", encode_sheets_data(json.dumps(body)))

    def _ops(self, node: str) -> list[dict]:
        return frappe.get_all(
            OP_LOG,
            filters={"sheet": self._docname(node)},
            fields=["seq", "op_type", "summary"],
            order_by="seq asc",
        )

    def _xlsx_node(self, raw: bytes, title="book.xlsx") -> str:
        from frappe.storage.blob import put_blob

        blob = put_blob(io.BytesIO(raw), is_private=True, filename=title)
        return create_file(
            self.admin, self.root.node, title, blob=blob.name, size=blob.file_size, mime=blob.mime_type
        )

    def _as(self, user: str):
        frappe.set_user(user)
        self.addCleanup(frappe.set_user, "Administrator")

    # creation, and the immutable link

    def test_a_new_sheet_carries_its_node_and_its_node_carries_it(self):
        node = self._sheet(title="Budget")
        row = frappe.db.get_value(
            "Drive Node", node, ("kind", "mime", "title", "content_doctype", "content_docname"), as_dict=True
        )
        self.assertEqual(row.kind, "document")
        self.assertEqual(row.mime, "frappe/sheet")
        self.assertEqual(row.title, "Budget")
        self.assertEqual(row.content_doctype, DOCTYPE)
        self.assertEqual(frappe.db.get_value(DOCTYPE, row.content_docname, "node"), node)

    def test_a_new_sheet_starts_empty_and_has_no_title_of_its_own(self):
        node = self._sheet(title="Budget")
        row = frappe.db.get_value(DOCTYPE, self._docname(node), ("title", "sheets_data"), as_dict=True)
        self.assertFalse(row.title, "Drive owns the title; §10.2 forbids a mirror")
        self.assertEqual(decode_sheets_data(row.sheets_data), "{}")

    def test_a_new_sheet_opens_its_history_with_a_create_op(self):
        node = self._sheet()
        self.assertEqual([op["op_type"] for op in self._ops(node)], ["create"])

    def test_the_head_seq_names_the_op_that_wrote_the_body(self):
        node = self._sheet()
        ops = self._ops(node)
        self.assertEqual(frappe.db.get_value(DOCTYPE, self._docname(node), "head_seq"), ops[-1]["seq"])

    def test_a_linked_sheet_cannot_write_the_frozen_legacy_title(self):
        node = self._sheet()
        sheet = frappe.get_doc(DOCTYPE, self._docname(node))
        sheet.title = "A mirror of the node title"
        with self.assertRaises(DriveConflict):
            sheet.save(ignore_permissions=True)

    def test_a_linked_sheet_cannot_write_the_frozen_trash_columns(self):
        node = self._sheet()
        sheet = frappe.get_doc(DOCTYPE, self._docname(node))
        sheet.trashed = 1
        with self.assertRaises(DriveConflict):
            sheet.save(ignore_permissions=True)

    def test_a_save_keeps_the_build_title_so_the_rollback_source_survives(self):
        """§14.11: the post-Build rollback is "ship the old code", which reads it."""
        node = self._sheet()
        docname = self._docname(node)
        frappe.db.set_value(DOCTYPE, docname, "title", "The Build title", update_modified=False)
        sheet = frappe.get_doc(DOCTYPE, docname)
        sheet.sheets_data = encode_sheets_data('{"a":1}')
        sheet.save(ignore_permissions=True)
        self.assertEqual(frappe.db.get_value(DOCTYPE, docname, "title"), "The Build title")

    def test_a_sheet_cannot_be_repointed_at_another_node(self):
        first = self._sheet(title="One")
        second = self._sheet(title="Two")
        sheet = frappe.get_doc(DOCTYPE, self._docname(first))
        sheet.node = second
        with self.assertRaises(DriveConflict):
            sheet.save(ignore_permissions=True)

    # copy

    def test_a_copy_carries_the_body(self):
        node = self._sheet(title="Source")
        self._write_body(node, {"sheet": {"v": 2, "current": "Data", "sheets": {}}})
        copy = drive.copy(node, self.root.node, title="Copy")
        self.assertEqual(json.loads(self._body(copy))["sheet"]["current"], "Data")

    def test_a_copy_gets_no_history_of_its_own_beyond_its_creation(self):
        """§8.9: a copy starts clean; it does not inherit the source's ops."""
        node = self._sheet(title="Source")
        sheets._append_op(self._docname(node), "save", "Saved")
        copy = drive.copy(node, self.root.node, title="Copy")
        self.assertEqual([op["op_type"] for op in self._ops(copy)], ["create"])

    def test_a_copy_gets_its_own_collaborative_document(self):
        """Inheriting the source's Y.Doc would rebase every editor onto the copy."""
        node = self._sheet(title="Source")
        frappe.get_doc(
            {"doctype": COLLAB_STATE, "sheet": self._docname(node), "ydoc_state": "AAA", "byte_size": 3}
        ).insert(ignore_permissions=True)
        copy = drive.copy(node, self.root.node, title="Copy")
        self.assertFalse(frappe.db.exists(COLLAB_STATE, self._docname(copy)))

    def test_a_copy_has_a_title_only_on_its_node(self):
        node = self._sheet(title="Source")
        copy = drive.copy(node, self.root.node, title="Copy")
        self.assertEqual(frappe.db.get_value("Drive Node", copy, "title"), "Copy")
        self.assertFalse(frappe.db.get_value(DOCTYPE, self._docname(copy), "title"))

    # xlsx import

    def test_an_xlsx_becomes_a_sheet_under_the_folder_the_caller_chose(self):
        source = self._xlsx_node(simple_workbook())
        node = drive.import_document(self.root.node, "Imported", content_doctype=DOCTYPE, from_node=source)
        row = frappe.db.get_value("Drive Node", node, ("kind", "mime", "title"), as_dict=True)
        self.assertEqual(row.kind, "document")
        self.assertEqual(row.mime, "frappe/sheet")
        self.assertEqual(row.title, "Imported")

    def test_an_import_carries_the_workbook_the_file_held(self):
        source = self._xlsx_node(simple_workbook())
        node = drive.import_document(self.root.node, "Imported", content_doctype=DOCTYPE, from_node=source)
        rows = json.loads(self._body(node))["sheet"]["sheets"]["Data"]["rows"]
        self.assertEqual(rows["0"], ["Name", "Amount"])

    def test_an_import_says_so_in_the_history(self):
        source = self._xlsx_node(simple_workbook())
        node = drive.import_document(self.root.node, "Imported", content_doctype=DOCTYPE, from_node=source)
        self.assertEqual([op["op_type"] for op in self._ops(node)], ["import"])

    def test_an_import_leaves_the_source_file_exactly_as_it_was(self):
        """An import is neither a move nor a copy."""
        source = self._xlsx_node(simple_workbook())
        before = frappe.db.get_value("Drive Node", source, ("parent", "state", "blob"), as_dict=True)
        drive.import_document(self.root.node, "Imported", content_doctype=DOCTYPE, from_node=source)
        after = frappe.db.get_value("Drive Node", source, ("parent", "state", "blob"), as_dict=True)
        self.assertEqual(before, after)

    def test_a_caller_who_cannot_read_the_file_cannot_import_it(self):
        source = self._xlsx_node(simple_workbook())
        self._as(OTHER)
        with self.assertRaises((DriveNotFound, DriveForbidden)):
            drive.import_document(self.other_root.node, "Imported", content_doctype=DOCTYPE, from_node=source)

    def test_a_file_that_is_not_a_workbook_is_refused_and_leaves_no_node(self):
        source = self._xlsx_node(b"not a spreadsheet", title="notes.txt")
        before = frappe.db.count("Drive Node")
        with self.assertRaises(Exception):
            drive.import_document(self.root.node, "Imported", content_doctype=DOCTYPE, from_node=source)
        self.assertEqual(frappe.db.count("Drive Node"), before)

    # versions

    def test_a_version_stores_the_workbook_and_the_head_seq(self):
        node = self._sheet()
        self._write_body(node, {"sheet": {"v": 2, "current": "Data", "sheets": {}}})
        seq = drive.take_version(node, kind="named", label="Before the edit")
        self.assertTrue(frappe.db.exists("Drive Node Version", {"node": node, "seq": seq}))

    def test_a_sheet_body_is_free_and_a_version_is_charged(self):
        """§7.1: the body is a column, so only the version blob is admitted."""
        node = self._sheet()
        self._write_body(node, {"sheet": {"v": 2, "current": "Data", "sheets": {}}})
        before = frappe.db.get_value("Drive Root", self.root.name, "used_bytes")
        drive.take_version(node)
        after = frappe.db.get_value("Drive Root", self.root.name, "used_bytes")
        self.assertGreater(after, before, "a version blob is charged")

        self._write_body(node, {"sheet": {"v": 2, "current": "Other", "sheets": {}}})
        self.assertEqual(
            frappe.db.get_value("Drive Root", self.root.name, "used_bytes"),
            after,
            "the body itself costs the root nothing",
        )

    def test_a_restore_puts_the_body_back(self):
        node = self._sheet()
        self._write_body(node, {"sheet": {"v": 2, "current": "First", "sheets": {}}})
        seq = drive.take_version(node, kind="named", label="First")
        self._write_body(node, {"sheet": {"v": 2, "current": "Second", "sheets": {}}})

        restore_version(self.admin, node, seq)
        self.assertEqual(json.loads(self._body(node))["sheet"]["current"], "First")

    def test_a_restore_is_one_op_of_its_own_at_a_fresh_seq(self):
        node = self._sheet()
        self._write_body(node, {"sheet": {"v": 2, "current": "First", "sheets": {}}})
        seq = drive.take_version(node, kind="named", label="First")
        self._write_body(node, {"sheet": {"v": 2, "current": "Second", "sheets": {}}})
        before = self._ops(node)[-1]["seq"]

        restore_version(self.admin, node, seq)
        ops = self._ops(node)
        self.assertEqual(ops[-1]["op_type"], "restore")
        self.assertGreater(ops[-1]["seq"], before, "head_seq never regresses")
        self.assertEqual(frappe.db.get_value(DOCTYPE, self._docname(node), "head_seq"), ops[-1]["seq"])

    def test_a_restore_takes_a_version_of_what_it_replaces(self):
        node = self._sheet()
        self._write_body(node, {"sheet": {"v": 2, "current": "First", "sheets": {}}})
        seq = drive.take_version(node, kind="named", label="First")
        self._write_body(node, {"sheet": {"v": 2, "current": "Second", "sheets": {}}})

        restore_version(self.admin, node, seq)
        self.assertGreater(frappe.db.count("Drive Node Version", {"node": node}), 1)

    # purge

    def test_a_purge_takes_the_sheet_and_every_app_owned_row_with_it(self):
        node = self._sheet()
        docname = self._docname(node)
        frappe.get_doc(
            {"doctype": COLLAB_STATE, "sheet": docname, "ydoc_state": "AAA", "byte_size": 3}
        ).insert(ignore_permissions=True)

        update(self.admin, node, state="Trashed")
        purge(self.admin, node)

        self.assertFalse(frappe.db.exists(DOCTYPE, docname))
        self.assertFalse(frappe.db.exists(OP_LOG, {"sheet": docname}))
        self.assertFalse(frappe.db.exists(COLLAB_STATE, docname))
        self.assertFalse(frappe.db.exists("Sheet Seq", docname))

    def test_a_purge_leaves_no_copy_of_the_workbook_behind(self):
        """`delete_permanently` is what makes a purge a purge."""
        node = self._sheet()
        docname = self._docname(node)
        update(self.admin, node, state="Trashed")
        purge(self.admin, node)
        self.assertFalse(frappe.db.exists("Deleted Document", {"deleted_name": docname}))

    # media discovery

    def test_the_sweep_is_told_every_id_a_body_still_names(self):
        node = self._sheet()
        media = self._media_node("picture.png")
        self._write_body(node, {"sheet": {"sheets": {"Data": {"rows": {"0": [media]}}}}})
        self.assertIn(media, sheets.used_nodes(self._docname(node)))

    def _media_node(self, title: str) -> str:
        from frappe.storage.blob import put_blob

        blob = put_blob(io.BytesIO(b"\x89PNG payload"), is_private=True, filename=title)
        with mock.patch("suite.drive._core.previews.enqueue_render"):
            return create_file(
                self.admin, self.root.node, title, blob=blob.name, size=blob.file_size, mime="image/png"
            )

    # satellites

    def test_the_op_log_takes_its_rights_from_the_sheet_s_node(self):
        node = self._sheet()
        op = frappe.get_all(OP_LOG, filters={"sheet": self._docname(node)}, pluck="name")[0]

        self._as(OTHER)
        self.assertFalse(satellite_has_permission(frappe.get_doc(OP_LOG, op), "read", OTHER))

    def test_a_reader_of_the_sheet_may_read_its_op_log(self):
        node = self._sheet()
        op = frappe.get_all(OP_LOG, filters={"sheet": self._docname(node)}, pluck="name")[0]
        grant(node, OTHER, drive.READ, self.admin)

        self._as(OTHER)
        self.assertTrue(satellite_has_permission(frappe.get_doc(OP_LOG, op), "read", OTHER))

    def test_the_op_log_list_is_scoped_to_readable_sheets(self):
        # Both sheets sit in USER's root, so the grant is the only difference
        # between them. A sheet in OTHER's own root proves nothing here: a
        # Personal root anchors MANAGE to its user, so OTHER reads all of it.
        mine = self._sheet(title="Mine")
        theirs = self._sheet(title="Theirs")
        grant(mine, OTHER, drive.READ, self.admin)

        self._as(OTHER)
        condition = satellite_query_conditions(OTHER, OP_LOG)
        # `get_list`, not `get_all`: `get_all` sets `ignore_permissions=True`,
        # so it never runs `permission_query_conditions` and would list both.
        listed = frappe.get_list(OP_LOG, pluck="sheet", limit_page_length=0)
        self.assertTrue(condition, "a non-admin gets a predicate")
        self.assertIn(self._docname(mine), listed)
        self.assertNotIn(self._docname(theirs), listed)
        self.assertEqual(self._ops_matching(condition), set(listed))

    def _ops_matching(self, predicate: str) -> set[str]:
        """The sheets the raw predicate admits, as `get_list` would apply it."""
        rows = frappe.db.sql(f"SELECT `sheet` FROM `tabSheet Op Log` WHERE {predicate}")
        return {row[0] for row in rows}

    def test_the_collaborative_document_takes_the_same_rights(self):
        node = self._sheet()
        state = frappe.get_doc(
            {"doctype": COLLAB_STATE, "sheet": self._docname(node), "ydoc_state": "AAA", "byte_size": 3}
        ).insert(ignore_permissions=True)

        self._as(OTHER)
        self.assertFalse(satellite_has_permission(state, "read", OTHER))

    def test_edit_is_what_changes_a_satellite_not_read(self):
        node = self._sheet()
        op = frappe.get_all(OP_LOG, filters={"sheet": self._docname(node)}, pluck="name")[0]
        grant(node, OTHER, drive.READ, self.admin)

        self._as(OTHER)
        document = frappe.get_doc(OP_LOG, op)
        self.assertTrue(satellite_has_permission(document, "read", OTHER))
        self.assertFalse(satellite_has_permission(document, "write", OTHER))


# ── the probe the site gate ran by hand ──────────────────────────────────────

GATE_OWNER = "sheets-gate-owner@example.com"
GATE_VICTIM = "sheets-gate-victim@example.com"
GATE_USERS = (GATE_OWNER, GATE_VICTIM)

# What a gate fixture writes and has to hand back. Rows in `_GATE_ADDED` are
# deleted outright after each test. Rows in `_GATE_WITNESS` are only counted,
# because Drive's own purge is what removes them and a raw delete here would
# hide a purge that never ran.
_GATE_ADDED = (
    OP_LOG,
    "Sheet Seq",
    "Sheet Snapshot",
    COLLAB_STATE,
    "Sheet Cell",
    "DocShare",
    "Notification Log",
    "Email Queue",
    "Error Log",
    "Version",
    "Comment",
    "ToDo",
    DOCTYPE,
)
# `File` and its activity log are witnesses rather than deletions: they belong
# to Drive, and a test may not write a `Drive *` table (ARCHITECTURE.md rule
# 2.2). `_drop_backing_file` hands them back through the File controller.
_GATE_WITNESS = (
    "Drive Root",
    "Drive Node",
    "Drive Grant",
    "File",
    "File Blob",
    "Drive Entity Activity Log",
    "User",
)


def _census(doctypes: tuple[str, ...]) -> dict[str, set[str]]:
    """Every name each table holds right now."""
    census = {}
    for doctype in doctypes:
        census[doctype] = set(frappe.get_all(doctype, pluck="name"))
    return census


def _open(docname: str) -> dict:
    """The stock row read, the way `/api/v2/document/Sheet/<name>` runs it.

    `frappe.get_doc` on its own checks nothing (`frappe/model/document.py:336`),
    so a probe built on it answers for every caller and proves nothing.
    `frappe.client.get` is `get_doc` plus `check_permission`
    (`frappe/client.py:107-113`), which is where the guard and the `DocShare`
    both get their say.
    """
    import frappe.client

    return frappe.client.get(DOCTYPE, docname)


def _listed(doctype: str, field: str = "name") -> list[str]:
    """`get_list`, which runs the permission hook and ORs the caller's shares.

    `get_all` sets `ignore_permissions`, so it never reaches either one.
    """
    return frappe.get_list(doctype, pluck=field, limit_page_length=0)


class TestTheGateProbes(IntegrationTestCase):
    """Gate step 6, as tests rather than a `bench console` script.

    Step 6 is the `DocShare` bypass. Frappe widens a denied row check with
    `false_if_not_shared` (`frappe/permissions.py:214-216`) and ORs the
    caller's shared names around a list predicate
    (`frappe/database/query.py:1737-1742`). Neither can be answered from inside
    a hook, so both guards refuse instead. The ticket called this unreachable
    by test. It is reachable: the sheet is created through Drive, the share is
    written the way Build inherited one, and the read runs through
    `frappe.client.get` and `frappe.get_list`.

    Every share here is hand-written, because ticket 29 refuses a new one
    (`refuse_governed_share`). A row Build inherited is the only kind the
    guards can ever meet, so it is the only kind the fixtures make.

    Gate step 9 was the legacy arm, on a sheet with no node. That state is
    gone: `require_node` holds §5.13 for a registered doctype, so the legacy
    endpoints have no row left to answer for.

    The two users, their roots, and every row a test writes are taken back off
    the site afterwards, and `_assert_no_residue` is what checks it.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.addClassCleanup(_drop_fixture_users, GATE_USERS)

    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        for user in GATE_USERS:
            ensure_user(user)
        _purge_fixture_roots(GATE_USERS)
        frappe.db.commit()
        self.census = _census(_GATE_ADDED + _GATE_WITNESS)
        # Registered first, so it runs last: the residue check has to see the
        # site after every other cleanup, not between two of them.
        self.addCleanup(self._assert_no_residue)
        self.addCleanup(self._drop_gate_rows)
        self.addCleanup(_purge_fixture_roots, GATE_USERS)
        self.addCleanup(frappe.set_user, "Administrator")

    # fixtures

    def _linked_sheet(self) -> tuple[str, str]:
        """One sheet linked the way Build links it: through Drive, then read back."""
        with activated():
            root = create_root(kind="Personal", title="Gate Root", user=GATE_OWNER)
            node = drive.create_document(root.node, "gate-linked", content_doctype=DOCTYPE)
        return node, frappe.db.get_value("Drive Node", node, "content_docname")

    def _share(self, doctype: str, name: str, user: str | None = None, **rights) -> None:
        """One `DocShare` the way a site carried it before adoption.

        `frappe.share.add` saves through `doc.save()` (`frappe/share.py:82`),
        so `refuse_governed_share` refuses it on a governed doctype.
        `ignore_validate` is what a Build-inherited row skipped, and it also
        keeps `cascade_permissions_downwards` off, so a read-only row stays
        read-only. `_drop_gate_rows` takes the row back.
        """
        frappe.set_user("Administrator")
        share = frappe.get_doc(
            {"doctype": "DocShare", "share_doctype": doctype, "share_name": name, "user": user, **rights}
        )
        share.flags.ignore_validate = True
        share.insert(ignore_permissions=True)

    def _as(self, user: str) -> None:
        frappe.set_user(user)
        self.addCleanup(frappe.set_user, "Administrator")

    # cleanup

    def _drop_gate_rows(self) -> None:
        frappe.set_user("Administrator")
        sheets = set(frappe.get_all(DOCTYPE, pluck="name")) - self.census[DOCTYPE]
        for doctype in _GATE_ADDED:
            added = tuple(set(frappe.get_all(doctype, pluck="name")) - self.census[doctype])
            if not added:
                continue
            frappe.db.delete(doctype, {"name": ("in", added)})
            if doctype == "Email Queue":
                frappe.db.delete("Email Queue Recipient", {"parent": ("in", added)})
        # The backing `File` last. Its `after_delete` cascades into the content
        # document (`suite/drive/overrides/file.py:146-152`), so it has to run
        # when there is no longer one to take with it.
        for sheet in sheets:
            _drop_backing_file(sheet)
        frappe.db.commit()

    def _assert_no_residue(self) -> None:
        frappe.set_user("Administrator")
        left = {}
        for doctype, before in self.census.items():
            added = sorted(set(frappe.get_all(doctype, pluck="name")) - before)
            if added:
                left[doctype] = added
        frappe.db.commit()
        self.assertEqual(left, {}, "a gate fixture left rows on the site")

    # ── gate step 6: the DocShare bypass, on the hooks the site ships ────────

    def test_a_named_share_cannot_open_a_linked_sheet(self):
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, GATE_VICTIM, read=1, write=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _open(docname)

    def test_a_named_share_cannot_list_a_linked_sheet(self):
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, GATE_VICTIM, read=1, write=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _listed(DOCTYPE)

    def test_a_named_share_does_not_widen_the_op_log_list(self):
        """The share is on the sheet, so the child list has nothing to OR.

        It answers rather than refusing, and what it answers must not name the
        linked sheet. The row count below is the control: it proves an empty
        answer is the guard, not an empty table.
        """
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, GATE_VICTIM, read=1)
        self.assertTrue(frappe.get_all(OP_LOG, filters={"sheet": docname}, pluck="name"))

        self._as(GATE_VICTIM)
        self.assertNotIn(docname, _listed(OP_LOG, "sheet"))

    def test_an_everyone_share_cannot_open_a_linked_sheet(self):
        """`get_shared` answers an `everyone` row for every user but a Guest
        (`frappe/share.py:188-190`), so it widens exactly as a named row does."""
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, None, read=1, everyone=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _open(docname)

    def test_an_everyone_share_cannot_list_a_linked_sheet(self):
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, None, read=1, everyone=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _listed(DOCTYPE)

    def test_an_everyone_share_does_not_widen_the_op_log_list(self):
        _node, docname = self._linked_sheet()
        self._share(DOCTYPE, docname, None, read=1, everyone=1)

        self._as(GATE_VICTIM)
        self.assertNotIn(docname, _listed(OP_LOG, "sheet"))

    def test_a_share_on_one_op_log_row_refuses_the_whole_op_log_list(self):
        """The one link no node column can scope: the share is on the child.

        `DriveForbidden`, not `frappe.PermissionError`: the guard is
        `drive.refuse_shared_child_rows`, and `frappe.desk.notifications` and
        `frappe.desk.desktop` swallow a `PermissionError`.
        """
        _node, docname = self._linked_sheet()
        op = frappe.get_all(OP_LOG, filters={"sheet": docname}, pluck="name")[0]
        self._share(OP_LOG, op, GATE_VICTIM, read=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _listed(OP_LOG, "sheet")

    def test_a_share_on_one_snapshot_row_refuses_the_whole_snapshot_list(self):
        """`Sheet Snapshot` keeps its own guard past activation, and gets the
        same refusal for the same reason."""
        _node, docname = self._linked_sheet()
        snapshot = self._preserved_snapshot(docname)
        self._share("Sheet Snapshot", snapshot, GATE_VICTIM, read=1)

        self._as(GATE_VICTIM)
        with self.assertRaises(DriveForbidden):
            _listed("Sheet Snapshot", "sheet")

    def test_a_new_snapshot_under_a_linked_sheet_is_refused_on_insert(self):
        """`SheetSnapshot.validate` is the guard `_preserved_snapshot` skips.

        Build preserves rows that predate the link, so the fixture needs
        `ignore_validate`. Nothing else may: a legacy snapshot written after
        the link would be history Drive does not own (§8, "Refuse linked-Sheet
        mutations in every legacy create, restore, label, pin, delete, and
        pruning path").
        """
        _node, docname = self._linked_sheet()
        row = frappe.get_doc(
            {"doctype": "Sheet Snapshot", "sheet": docname, "seq": 2, "kind": "auto", "sheets_data": "{}"}
        )
        with self.assertRaises(frappe.ValidationError):
            row.insert(ignore_permissions=True)

    def test_deleting_a_preserved_snapshot_under_a_linked_sheet_is_refused(self):
        """`SheetSnapshot.on_trash` is the other half. Drive prunes migrated
        history; the legacy pruner must not reach a linked sheet's rows."""
        _node, docname = self._linked_sheet()
        snapshot = self._preserved_snapshot(docname)
        with self.assertRaises(frappe.ValidationError):
            frappe.delete_doc("Sheet Snapshot", snapshot, ignore_permissions=True)

    def _preserved_snapshot(self, docname: str) -> str:
        """One `Sheet Snapshot` row standing in for one Build preserved.

        It predates the link, so it never met `SheetSnapshot.validate`, which
        refuses a linked parent.
        """
        row = frappe.get_doc(
            {"doctype": "Sheet Snapshot", "sheet": docname, "seq": 1, "kind": "auto", "sheets_data": "{}"}
        )
        row.flags.ignore_validate = True
        return row.insert(ignore_permissions=True).name

    # the one share that cannot be written at all

    def test_a_new_share_on_a_linked_sheet_is_refused_outright(self):
        """`refuse_governed_share` is why a share on a linked sheet can only be
        one Build inherited, never one written after ticket 29. Through
        `frappe.share.add`, because that is the call Desk assignment makes."""
        import frappe.share

        _node, docname = self._linked_sheet()

        with self.assertRaises(DriveForbidden):
            frappe.share.add(DOCTYPE, docname, GATE_VICTIM, read=1, notify=False)


def _drop_fixture_users(users: tuple[str, ...] = (USER, OTHER)) -> None:
    """Take the fixture users, and the roots their inserts provisioned, back off.

    Every class here creates its users in `setUpClass`, so leaving them behind
    grows the site by one account and one Personal root per module that runs.
    """
    frappe.set_user("Administrator")
    _purge_fixture_roots(users)
    for user in users:
        frappe.delete_doc("User", user, force=1, ignore_permissions=True, ignore_missing=True)
    frappe.db.commit()


def _drop_backing_file(docname: str) -> None:
    """Take the legacy `File` a `Sheet` insert provisions off the site.

    `File.permanent_delete` marks the row `Removed` rather than deleting it
    (`suite/drive/overrides/file.py:402-406`), so a fixture that deletes only
    its `Sheet` leaves one orphan row behind on every run. Ticket 23 removes
    the backing for good; until then a test takes back what it made.

    Through `delete_doc`, not a raw delete: `File.after_delete` is what clears
    the activity log and the favourites that hang off the row
    (`suite/drive/overrides/file.py:131-145`), and a test may not write a
    `Drive *` table itself (ARCHITECTURE.md rule 2.2).
    """
    for name in frappe.get_all(
        "File", filters={"content_doctype": DOCTYPE, "content_docname": docname}, pluck="name"
    ):
        frappe.delete_doc("File", name, force=1, ignore_permissions=True, ignore_missing=True)


def _purge_fixture_roots(users: tuple[str, ...] = (USER, OTHER)) -> None:
    """Hand back every Drive root the named fixture users own, through Drive's purge.

    Every user this module names belongs to it alone, so the filter can never
    reach a live account or another test. A `User` insert provisions a Personal
    root of its own, so this also runs before a fixture creates one.
    """
    admin = Principals("Administrator", ("Administrator",), (), is_admin=True)
    roots = frappe.get_all("Drive Root", filters={"user": ["in", users]}, pluck="name")
    # Purging a document node calls the app's `on_purge`, which Drive reads from
    # the registry, so the purge runs against a cache built for this request.
    with activated():
        for root in roots:
            if frappe.db.get_value("Drive Root", root, "state") == "Active":
                update_root(root, admin, state="Archived")
            purge_root(root, admin)


if __name__ == "__main__":
    unittest.main()
